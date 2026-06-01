from __future__ import annotations

import csv
import unittest
from pathlib import Path
from sys import path as sys_path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

sys_path.insert(0, str(Path(__file__).resolve().parents[1]))

from scripts.wp_excel_to_woocommerce_csv import convert_workbook


class ConvertWorkbookTest(unittest.TestCase):
    def test_creates_products_csv_and_manifest(self) -> None:
        repo_root = Path(__file__).resolve().parents[1]
        output_dir = repo_root / "output"
        input_path = output_dir / "__test_products.xlsx"
        output_path = output_dir / "__test_products.csv"
        manifest_path = output_dir / "__test_downloads_manifest.csv"
        asset_path = output_dir / "__test_guide.pdf"

        for path in [input_path, output_path, manifest_path, asset_path]:
            try:
                if path.exists():
                    path.unlink()
            except PermissionError:
                pass

        try:
            asset_path.write_text("pdf", encoding="utf-8")

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Products"
            sheet.append(
                [
                    "Название",
                    "Артикул",
                    "Цена",
                    "Краткое описание",
                    "Описание",
                    "Категории",
                    "Теги",
                    "Файл",
                    "SEO title",
                    "SEO description",
                    "Ключевая фраза",
                ]
            )
            sheet.append(
                [
                    "Digital Kit",
                    "DK-001",
                    "29.00",
                    "Short text",
                    "Long text",
                    "Art, Digital",
                    "pack|pdf",
                    asset_path.name,
                    "SEO title",
                    "SEO description",
                    "digital kit",
                ]
            )
            workbook.save(input_path)

            products, missing_downloads = convert_workbook(
                input_path=input_path,
                output_path=output_path,
                manifest_path=manifest_path,
                sheet_name="Products",
                download_base_url="https://example.com/downloads",
            )

            self.assertEqual(products, 1)
            self.assertEqual(missing_downloads, 0)
            self.assertTrue(output_path.exists())
            self.assertTrue(manifest_path.exists())

            with output_path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(rows[0]["Name"], "Digital Kit")
            self.assertEqual(rows[0]["Slug"], "digital-kit")
            self.assertEqual(rows[0]["Regular price"], "29.00")
            self.assertEqual(rows[0]["Categories"], "Art | Digital")
            self.assertEqual(rows[0]["Tags"], "pack | pdf")
            self.assertEqual(rows[0]["Download 1 URL"], "https://example.com/downloads/__test_guide.pdf")
            self.assertEqual(rows[0]["Meta: _yoast_wpseo_title"], "SEO title")

            with manifest_path.open("r", encoding="utf-8-sig", newline="") as handle:
                manifest_rows = list(csv.DictReader(handle))
            self.assertEqual(manifest_rows[0]["Exists"], "yes")
            self.assertIn("__test_guide.pdf", manifest_rows[0]["Resolved path"])
        finally:
            for path in [input_path, output_path, manifest_path, asset_path]:
                try:
                    if path.exists():
                        path.unlink()
                except PermissionError:
                    pass

    def test_uses_product_folder_images_and_zip(self) -> None:
        repo_root = Path(__file__).resolve().parents[1]
        output_dir = repo_root / "output"
        input_path = output_dir / "__folder_products.xlsx"
        output_path = output_dir / "__folder_products.csv"
        manifest_path = output_dir / "__folder_downloads_manifest.csv"
        image_path = output_dir / "__folder_01.png"
        zip_path = output_dir / "__folder_product.zip"

        for path in [input_path, output_path, manifest_path, image_path, zip_path]:
            try:
                if path.exists():
                    path.unlink()
            except PermissionError:
                pass

        try:
            image_path.write_bytes(b"png")
            zip_path.write_bytes(b"zip")

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Products"
            sheet.append(["title", "sku", "regular_price", "product_folder"])
            sheet.append(["Folder Product", "FP-001", "31.00", "."])
            workbook.save(input_path)

            products, missing_downloads = convert_workbook(
                input_path=input_path,
                output_path=output_path,
                manifest_path=manifest_path,
                sheet_name="Products",
                download_base_url="https://example.com/media",
            )

            self.assertEqual(products, 1)
            self.assertEqual(missing_downloads, 0)

            with output_path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertIn("https://example.com/media/__folder_01.png", rows[0]["Images"])
            self.assertEqual(rows[0]["Download 1 URL"], "https://example.com/media/__folder_product.zip")

            with manifest_path.open("r", encoding="utf-8-sig", newline="") as handle:
                manifest_rows = list(csv.DictReader(handle))
            kinds = {row["Kind"] for row in manifest_rows}
            self.assertIn("folder", kinds)
            self.assertIn("image", kinds)
            self.assertIn("download", kinds)
        finally:
            for path in [input_path, output_path, manifest_path, image_path, zip_path]:
                try:
                    if path.exists():
                        path.unlink()
                except PermissionError:
                    pass

    def test_supports_user_workbook_columns_and_updates_status_date(self) -> None:
        repo_root = Path(__file__).resolve().parents[1]
        output_dir = repo_root / "output"
        input_path = output_dir / "__user_columns.xlsx"
        output_path = output_dir / "__user_columns.csv"
        manifest_path = output_dir / "__user_columns_manifest.csv"
        image_path = output_dir / "__user_main.jpg"
        zip_path = output_dir / "__user_product.zip"

        for path in [input_path, output_path, manifest_path, image_path, zip_path]:
            try:
                if path.exists():
                    path.unlink()
            except PermissionError:
                pass

        try:
            image_path.write_bytes(b"jpg")
            zip_path.write_bytes(b"zip")

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Etsy Listings"
            sheet.append(
                [
                    "File Name",
                    "Photos",
                    "Title",
                    "Description",
                    "Price",
                    "Alt Text для Google",
                    "Tags",
                    "Digital File (ZIP)",
                    "Путь к файлам",
                    "status",
                    "category",
                    "error_message",
                ]
            )
            sheet.append(
                [
                    "User Product",
                    image_path.name,
                    "User Product Title",
                    "Description",
                    "$3.83",
                    "Alt text",
                    "tag one, tag two",
                    zip_path.name,
                    str(output_dir),
                    "ERROR",
                    "Autumn",
                    "old error",
                ]
            )
            workbook.save(input_path)

            products, missing_downloads = convert_workbook(
                input_path=input_path,
                output_path=output_path,
                manifest_path=manifest_path,
                sheet_name="Etsy Listings",
                download_base_url="https://example.com/media",
                publication_status="draft",
                update_source_status=True,
            )

            self.assertEqual(products, 1)
            self.assertEqual(missing_downloads, 0)

            with output_path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(rows[0]["Published"], "0")
            self.assertEqual(rows[0]["Regular price"], "3.83")
            self.assertEqual(rows[0]["Download 1 URL"], "https://example.com/media/__user_product.zip")
            self.assertEqual(rows[0]["Images"], "https://example.com/media/__user_main.jpg")

            updated = load_workbook(input_path)
            updated_sheet = updated["Etsy Listings"]
            headers = [updated_sheet.cell(1, column).value for column in range(1, updated_sheet.max_column + 1)]
            status_col = headers.index("status") + 1
            date_col = headers.index("status_date") + 1
            mode_col = headers.index("publish_mode") + 1
            self.assertEqual(date_col, status_col + 1)
            self.assertEqual(updated_sheet.cell(2, status_col).value, "CSV_READY")
            self.assertTrue(updated_sheet.cell(2, date_col).value)
            self.assertEqual(updated_sheet.cell(2, mode_col).value, "draft")
        finally:
            for path in [input_path, output_path, manifest_path, image_path, zip_path]:
                try:
                    if path.exists():
                        path.unlink()
                except PermissionError:
                    pass

    def test_status_update_falls_back_to_output_copy_when_source_is_locked(self) -> None:
        repo_root = Path(__file__).resolve().parents[1]
        output_dir = repo_root / "output"
        input_path = output_dir / "__locked_source.xlsx"
        output_path = output_dir / "__locked_source.csv"
        manifest_path = output_dir / "__locked_source_manifest.csv"
        fallback_path = output_dir / "__locked_source_with_status.xlsx"

        for path in [input_path, output_path, manifest_path, fallback_path]:
            try:
                if path.exists():
                    path.unlink()
            except PermissionError:
                pass

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["title", "status"])
        sheet.append(["Locked Source Product", ""])
        workbook.save(input_path)

        original_save = OpenpyxlWorkbook.save

        def save_with_locked_source(self, filename):
            if Path(filename) == input_path:
                raise PermissionError("locked")
            return original_save(self, filename)

        try:
            with patch.object(OpenpyxlWorkbook, "save", save_with_locked_source):
                result = convert_workbook(
                    input_path=input_path,
                    output_path=output_path,
                    manifest_path=manifest_path,
                    update_source_status=True,
                )

            self.assertEqual(result.product_count, 1)
            self.assertFalse(result.status_saved_to_source)
            self.assertEqual(result.status_workbook_path, fallback_path)
            self.assertTrue(fallback_path.exists())

            updated = load_workbook(fallback_path)
            updated_sheet = updated.active
            headers = [updated_sheet.cell(1, column).value for column in range(1, updated_sheet.max_column + 1)]
            status_col = headers.index("status") + 1
            self.assertEqual(updated_sheet.cell(2, status_col).value, "CSV_READY")
        finally:
            for path in [input_path, output_path, manifest_path, fallback_path]:
                try:
                    if path.exists():
                        path.unlink()
                except PermissionError:
                    pass


if __name__ == "__main__":
    unittest.main()
